import streamlit as st
import re
import io
import csv
import pandas as pd
from groq import Groq
from pypdf import PdfReader
from docx import Document as DocxDocument
import openpyxl
from pptx import Presentation

# ==========================================
# 1. SECURITY CHECK
# ==========================================
def get_real_key():
    p1 = "M."
    p2 = "LATIF"
    p3 = "PPC"
    return p1 + p2 + p3

if "authenticated" not in st.session_state:
    st.session_state.authenticated = False

if not st.session_state.authenticated:
    st.markdown("<h1 style='text-align: center;'>🔒 Secure AI Office Assistant</h1>", unsafe_allow_html=True)
    with st.form("login"):
        user_key = st.text_input("Enter Access Key:", type="password")
        if st.form_submit_button("Unlock"):
            if user_key == get_real_key():
                st.session_state.authenticated = True
                st.rerun()
            else:
                st.error("Invalid Key!")
    st.stop()

# ==========================================
# 2. SETUP AI & SESSION STATE
# ==========================================
client = Groq(api_key=st.secrets["GROQ_API_KEY"])
if "messages" not in st.session_state:
    st.session_state.messages = []
if "last_answer" not in st.session_state:
    st.session_state.last_answer = ""

st.title("🤖 AI Office Assistant (Data Analyst Edition)")

# ==========================================
# 3. CREATE TABS
# ==========================================
tab1, tab2 = st.tabs(["💬 Chat with AI (Multiple Files)", "📊 Auto Dashboard & AI Insights"])

# ==========================================
# TAB 1: AI CHATBOT (MULTI-FILE SUPPORT)
# ==========================================
with tab1:
    uploaded_files = st.file_uploader("Upload Documents (You can select multiple!)", type=["pdf", "docx", "xlsx", "pptx", "csv"], accept_multiple_files=True, key="chat")
    
    def read_pdf(file):
        return "\n".join([page.extract_text() for page in PdfReader(file).pages])
    def read_word(file):
        doc = DocxDocument(file)
        text = "\n".join([p.text for p in doc.paragraphs])
        for t in doc.tables:
            for r in t.rows: text += "\n" + " | ".join([c.text for c in r.cells])
        return text
    def read_pptx(file):
        text = ""
        for s in Presentation(file).slides:
            for sh in s.shapes:
                if sh.has_text_frame: text += sh.text_frame.text + "\n"
        return text
    def read_csv_txt(file):
        try: f = io.StringIO(file.read().decode('utf-8'))
        except: f = io.StringIO(file.read().decode('latin-1'))
        return "\n".join([",".join(row) for row in csv.reader(f)])
    def read_excel_txt(file):
        wb = openpyxl.load_workbook(file)
        text = ""
        for sn in wb.sheetnames:
            text += f"--- {sn} ---\n"
            for row in wb[sn].iter_rows(values_only=True): text += " | ".join([str(c) if c else "" for c in row]) + "\n"
        return text

    def chunk_text(text, size=4000, overlap=400):
        chunks, start = [], 0
        while start < len(text): chunks.append(text[start:start+size]); start += size - overlap
        return chunks
    def find_best_chunks(chunks, q):
        words = set(re.findall(r'\b\w+\b', q.lower())) - {"what","is","the","a","an","in","on","to","for","of","with","how","does","do","did","are","was","were","be","this","that","it","from","by","and","or","but","if"}
        if not words: return "\n".join(chunks[:3])
        scores = sorted([(sum(1 for w in words if w in c.lower()), c) for c in chunks], key=lambda x: x[0], reverse=True)
        return "\n".join([c for s, c in scores[:3]])

    if uploaded_files:
        all_text = ""
        with st.spinner(f"Reading {len(uploaded_files)} files..."):
            for file in uploaded_files:
                ft = file.name.split(".")[-1].lower()
                if ft=="pdf": all_text += read_pdf(file) + "\n\n"
                elif ft=="docx": all_text += read_word(file) + "\n\n"
                elif ft=="pptx": all_text += read_pptx(file) + "\n\n"
                elif ft=="csv": all_text += read_csv_txt(file) + "\n\n"
                elif ft=="xlsx": all_text += read_excel_txt(file) + "\n\n"
            chunks = chunk_text(all_text)
        st.success(f"✅ {len(uploaded_files)} files combined and indexed!")

        for msg in st.session_state.messages:
            with st.chat_message(msg["role"]): st.markdown(msg["content"])

        q = st.chat_input("Ask a question across ALL uploaded files...")
        if q:
            st.session_state.messages.append({"role": "user", "content": q})
            with st.chat_message("user"): st.markdown(q)
            with st.chat_message("assistant"):
                with st.spinner("Searching all documents..."):
                    ctx = find_best_chunks(chunks, q)
                    res = client.chat.completions.create(model="llama-3.3-70b-versatile", messages=[{"role":"user","content":f"Answer ONLY from context:\n{ctx}\n\nQ:{q}"}])
                    ans = res.choices[0].message.content
                    st.markdown(ans)
            st.session_state.messages.append({"role": "assistant", "content": ans})
            st.session_state.last_answer = ans

        if st.session_state.last_answer:
            if st.button("📥 Download Answer as Word"):
                doc = DocxDocument()
                doc.add_paragraph(st.session_state.last_answer)
                buf = io.BytesIO(); doc.save(buf); buf.seek(0)
                st.download_button("Save File", buf, "Answer.docx")

# ==========================================
# TAB 2: AUTO DASHBOARD + AI INSIGHTS
# ==========================================
with tab2:
    st.write("Upload an **Excel** or **CSV** file to visualize data and generate AI insights.")
    dash_file = st.file_uploader("Upload Data File", type=["xlsx", "csv"], key="dash")
    
    if dash_file:
        # Read data using Pandas
        if dash_file.name.endswith(".csv"):
            df = pd.read_csv(dash_file)
        else:
            df = pd.read_excel(dash_file)
            
        st.success("✅ Data loaded successfully!")
        
        with st.expander("View Raw Data"):
            st.dataframe(df.head(10))

        # --- CHART SECTION ---
        cols = df.columns.tolist()
        if len(cols) >= 2:
            c1, c2 = st.columns(2)
            with c1:
                x_axis = st.selectbox("X-Axis (Categories)", cols)
            with c2:
                y_axis = st.selectbox("Y-Axis (Numbers)", cols, index=1)
            
            chart_type = st.radio("Select Chart Type", ["Bar Chart", "Line Chart", "Area Chart"], horizontal=True)
            
            if st.button("📊 Generate Chart", use_container_width=True):
                try:
                    if chart_type == "Bar Chart":
                        st.bar_chart(df, x=x_axis, y=y_axis)
                    elif chart_type == "Line Chart":
                        st.line_chart(df, x=x_axis, y=y_axis)
                    elif chart_type == "Area Chart":
                        st.area_chart(df, x=x_axis, y=y_axis)
                except Exception as e:
                    st.error(f"Could not draw chart. Make sure Y-Axis has numbers. (Error: {e})")
        
        # --- NEW: AI INSIGHTS SECTION ---
        st.markdown("---")
        st.subheader("🧠 AI-Powered Data Insights")
        st.write("Let the AI analyze the math, find trends, and explain what the numbers mean.")
        
        if st.button("⚡ Generate AI Insights", use_container_width=True, type="primary"):
            with st.spinner("AI is analyzing statistical data..."):
                # 1. Get Mathematical Summary (Averages, Max, Min) - Safe for large files
                math_summary = df.describe(include='all').to_string()
                
                # 2. Get first few rows for context
                sample_data = df.head(15).to_string()
                
                # 3. Ask AI to act as a Data Analyst
                prompt = f"""You are a Senior Data Analyst. Analyze this dataset and provide 3-5 key business insights, trends, or anomalies. 
                
                Mathematical Summary (Mean, Min, Max, etc):
                {math_summary}
                
                Sample of the Data (First 15 rows):
                {sample_data}
                
                Please provide clear, actionable insights in bullet points. Do not just repeat the numbers, explain WHAT THEY MEAN for the business."""
                
                try:
                    res = client.chat.completions.create(
                        model="llama-3.3-70b-versatile", 
                        messages=[{"role":"user","content": prompt}]
                    )
                    st.success(res.choices[0].message.content)
                except Exception as e:
                    st.error(f"AI could not read this data format. Error: {e}")

# ==========================================
# SIDEBAR
# ==========================================
with st.sidebar:
    if st.button("🗑️ Clear Chat History"): st.session_state.messages = []; st.rerun()
    if st.button("🔒 Logout"): st.session_state.authenticated = False; st.rerun()
